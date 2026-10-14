"""Gera as tabelas LMS embutidas no AnthroCore a partir das fontes públicas da OMS (design D4 de port-anthro-engine).

Uso, da raiz do projeto:
    uv run --with openpyxl python scripts/generate_who_tables.py

Entradas (reference/who-tables/):
    {wfa,lhfa,bfa}-{boys,girls}-zscore-expanded-tables.xlsx  padrões 2006, dias 0 a 1856
    {wfa,hfa,bfa}-who2007-anthroplus-r.csv                   referência 2007, meses 60 a 229 (peso até 121)
    {wfa,lhfa,bfa}-anthro-r.csv                              pacote R anthro, dias 0 a 1826 (conferência)

Saída (Sources/AnthroCore/Resources/who/):
    std-{wfa,lhfa,bfa}.csv e ref-{wfa,lhfa,bfa}.csv com colunas sex,age,l,m,s (sex 1 = menino, 2 = menina).

Falha (código 1) se as planilhas e o pacote R anthro divergirem em algum dia de 0 a 1826, ou se faltar
alguma idade nos intervalos declarados.
"""

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCES = ROOT / "reference" / "who-tables"
OUTPUT = ROOT / "Sources" / "AnthroCore" / "Resources" / "who"

STANDARDS = {"wfa": "wfa", "lhfa": "lhfa", "bfa": "bfa"}
REFERENCE = {"wfa": "wfa", "lhfa": "hfa", "bfa": "bfa"}
REFERENCE_LAST_MONTH = {"wfa": 121, "lhfa": 229, "bfa": 229}
SEXES = {"boys": "1", "girls": "2"}


def text(value):
    """Valor numérico como texto, na forma mais curta que volta ao mesmo float (o que está na planilha)."""
    if isinstance(value, str):
        return value.strip()
    return repr(float(value)) if not float(value).is_integer() else str(int(value))


def read_expanded(indicator):
    rows = []
    for sex_name, sex in SEXES.items():
        sheet = load_workbook(SOURCES / f"{indicator}-{sex_name}-zscore-expanded-tables.xlsx", read_only=True).active
        values = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() for cell in next(values)]
        day, l, m, s = (header.index(name) for name in ("Day", "L", "M", "S"))
        for row in values:
            if row[day] is None:
                continue
            rows.append((sex, int(row[day]), text(row[l]), text(row[m]), text(row[s])))
    return rows


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as handle:
        return [(r["sex"], int(r["age"]), r["l"], r["m"], r["s"]) for r in csv.DictReader(handle)]


def check_contiguous(rows, first, last, label):
    for sex in SEXES.values():
        ages = {age for row_sex, age, *_ in rows if row_sex == sex}
        missing = [age for age in range(first, last + 1) if age not in ages]
        if missing:
            sys.exit(f"erro: {label}, sexo {sex}: faltam idades {missing[:5]}")


def check_against_anthro(rows, indicator):
    anthro = {(sex, age): (float(l), float(m), float(s)) for sex, age, l, m, s in read_csv(SOURCES / f"{indicator}-anthro-r.csv")}
    for sex, age, l, m, s in rows:
        if age > 1826:
            continue
        expected = anthro.get((sex, age))
        if expected != (float(l), float(m), float(s)):
            sys.exit(f"erro: {indicator} sexo {sex} dia {age}: planilha {(l, m, s)} difere do pacote R anthro {expected}")


def write(rows, path):
    rows = sorted(rows, key=lambda row: (row[0], row[1]))
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["sex", "age", "l", "m", "s"])
        writer.writerows(rows)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for indicator, source in STANDARDS.items():
        rows = read_expanded(source)
        check_contiguous(rows, 0, 1856, f"padrões {indicator}")
        check_against_anthro(rows, indicator)
        write(rows, OUTPUT / f"std-{indicator}.csv")
    for indicator, source in REFERENCE.items():
        rows = read_csv(SOURCES / f"{source}-who2007-anthroplus-r.csv")
        check_contiguous(rows, 60, REFERENCE_LAST_MONTH[indicator], f"referência {indicator}")
        write(rows, OUTPUT / f"ref-{indicator}.csv")
    print(f"tabelas gravadas em {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
